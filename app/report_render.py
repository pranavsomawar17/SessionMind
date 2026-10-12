"""Report layout: normalising model output, mind-map building, Word (.docx) and HTML rendering.

Kept free of FastAPI/Groq imports so it can be tested on its own.
"""
import difflib
import html
import re
from datetime import datetime
from pathlib import Path

MAX_TREE_DEPTH = 4
MAX_TREE_CHILDREN = 15


def H(x):
    return html.escape(str(x or ''))


def clock(sec):
    """Meeting-relative time as mm:ss (or h:mm:ss)."""
    try:
        sec = max(0, int(float(sec)))
    except Exception:
        return str(sec or '')
    h, rem = divmod(sec, 3600)
    mi, s_ = divmod(rem, 60)
    return f'{h}:{mi:02d}:{s_:02d}' if h else f'{mi:02d}:{s_:02d}'


def nice_dt(v):
    """ISO timestamp -> '12 Oct 2026, 18:07' in the server's local time zone."""
    try:
        d = datetime.fromisoformat(str(v))
        if d.tzinfo is not None:
            d = d.astimezone()
        return d.strftime('%d %b %Y, %H:%M')
    except Exception:
        return str(v or '')


def _clip(s, n=160):
    s = ' '.join(str(s or '').split())
    return s if len(s) <= n else s[:n - 1].rstrip() + '…'


# ---------------------------------------------------------------- normalisers
# The model sometimes returns plain strings or different key names. Every list the
# report shows goes through these, so an item is never dropped just for its shape.

def _first(d, *keys):
    for k in keys:
        v = d.get(k)
        if v not in (None, '', []):
            return str(v).strip()
    return ''


def clean_question(q, default_status='Unanswered'):
    if isinstance(q, str):
        q = {'question': q}
    if not isinstance(q, dict):
        return None
    text = _first(q, 'question', 'text', 'description', 'item')
    if not text:
        return None
    return {'question': text, 'status': _first(q, 'status') or default_status,
            'raised_by': _first(q, 'raised_by', 'speaker', 'raised'),
            'timestamp': _first(q, 'timestamp', 'time'), 'evidence': _first(q, 'evidence', 'quote')}


def clean_decision(d):
    if isinstance(d, str):
        d = {'decision': d}
    if not isinstance(d, dict):
        return None
    text = _first(d, 'decision', 'text', 'description', 'item')
    if not text:
        return None
    return {'decision': text, 'evidence': _first(d, 'evidence', 'quote'), 'timestamp': _first(d, 'timestamp', 'time')}


def clean_action(a):
    if isinstance(a, str):
        a = {'task': a}
    if not isinstance(a, dict):
        return None
    task = _first(a, 'task', 'action', 'description', 'text', 'item')
    if not task:
        return None
    return {'owner': _first(a, 'owner', 'assignee', 'who') or 'Unknown', 'task': task,
            'deadline': _first(a, 'deadline', 'due', 'due_date') or 'Unknown',
            'status': _first(a, 'status') or 'Assigned', 'evidence': _first(a, 'evidence', 'quote'),
            'timestamp': _first(a, 'timestamp', 'time')}


def _norm(s):
    return re.sub(r'\W+', ' ', str(s or '').lower()).strip()


def same_text(a, b):
    """Duplicate check. Near-identical wording counts as the same item only when the
    numbers in it match too, so 'item 5?' and 'item 10?' stay separate."""
    na, nb = _norm(a), _norm(b)
    if not na or not nb:
        return False
    if na == nb:
        return True
    if re.findall(r'\d+', na) != re.findall(r'\d+', nb):
        return False
    return difflib.SequenceMatcher(None, na, nb).ratio() >= 0.9


def merge_items(key, *groups):
    """Concatenate item lists, dropping later duplicates of `key`."""
    out = []
    for g in groups:
        for it in g:
            if it and not any(same_text(it[key], o[key]) for o in out):
                out.append(it)
    return out


def question_rows(rep):
    return [q for q in (clean_question(x) for x in (rep.get('open_questions') or [])) if q]


# ---------------------------------------------------------------- mind map

def _clean_node(n, depth=1):
    if isinstance(n, dict):
        label = _clip(n.get('label') or n.get('name') or n.get('text'), 120)
        raw_children = n.get('children') or []
    else:
        label = _clip(n, 120)
        raw_children = []
    if not label:
        return None
    kids = []
    if depth < MAX_TREE_DEPTH and isinstance(raw_children, list):
        for c in raw_children[:MAX_TREE_CHILDREN]:
            k = _clean_node(c, depth + 1)
            if k:
                kids.append(k)
    return {'label': label, 'children': kids}


def _fallback_tree(rep):
    """Used when the model returned no usable mind_map (or for older reports)."""
    branches = []
    for t in (rep.get('topics') or [])[:12]:
        if not isinstance(t, dict) or not t.get('name'):
            continue
        kids = [{'label': _clip(t['summary'], 160), 'children': []}] if t.get('summary') else []
        branches.append({'label': _clip(t['name'], 120), 'children': kids})
    decisions = [d['decision'] for d in (clean_decision(x) for x in (rep.get('decisions') or [])) if d]
    actions = [f"{a['owner']}: {a['task']}" for a in (clean_action(x) for x in (rep.get('action_items') or [])) if a]
    questions = [q['question'] for q in question_rows(rep)]
    risks = [str(r.get('risk') if isinstance(r, dict) else r) for r in (rep.get('risks') or []) if r]
    for name, items in (('Decisions', decisions), ('Action items', actions), ('Open questions', questions), ('Risks / blockers', risks)):
        if items:
            branches.append({'label': name, 'children': [{'label': _clip(i, 160), 'children': []} for i in items]})
    return {'label': _clip(rep.get('title') or 'Meeting', 120), 'children': branches}


def build_mind_map(rep):
    """Tree {'label', 'children'}: root = meeting, branches = themes, children = points under each theme."""
    mm = rep.get('mind_map')
    if isinstance(mm, dict):
        branches = [n for n in (_clean_node(b, 1) for b in (mm.get('branches') or [])[:12]) if n]
        if branches:
            return {'label': _clip(mm.get('root') or rep.get('title') or 'Meeting', 120), 'children': branches}
    return _fallback_tree(rep)


def tree_levels(node):
    kids = node.get('children') or []
    return 0 if not kids else 1 + max(tree_levels(k) for k in kids)


def mind_view(tree):
    """'tree' when the content has real hierarchy (2+ levels), otherwise 'points' (a key-point list)."""
    return 'tree' if tree_levels(tree) >= 2 else 'points'


def tree_lines(tree):
    lines = [tree['label']]

    def walk(node, prefix):
        kids = node.get('children') or []
        for i, k in enumerate(kids):
            last = i == len(kids) - 1
            lines.append(prefix + ('└── ' if last else '├── ') + k['label'])
            walk(k, prefix + ('    ' if last else '│   '))

    walk(tree, '')
    return lines


def key_points(tree):
    """[(theme, [point, ...])] for the key-points view."""
    return [(b['label'], [c['label'] for c in b.get('children', [])]) for b in tree.get('children', [])]


# ---------------------------------------------------------------- Word (.docx)

def render_docx(rep, transcript, screenshots, path, base_dir, tree=None):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    tree = tree or build_mind_map(rep)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)

    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    normal.paragraph_format.space_after = Pt(4)
    for name, size in (('Heading 1', 16), ('Heading 2', 13), ('Title', 22)):
        st = doc.styles[name]
        st.font.name = 'Calibri'
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)

    fp = sec.footer.paragraphs[0]
    fp.add_run('SessionMind report · page ').font.size = Pt(8)
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = '1'
    r.append(t)
    fld.append(r)
    fp._p.append(fld)

    GREY = (0x66, 0x66, 0x66)

    def para(text='', bold=False, italic=False, size=None, color=None):
        p = doc.add_paragraph()
        run = p.add_run(text or '')
        run.bold, run.italic = bold, italic
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor(*color)
        return p

    def bullets(items, empty='None recorded.'):
        items = [i for i in items if i not in (None, '')]
        if not items:
            para(empty, italic=True, color=GREY)
        for i in items:
            doc.add_paragraph(str(i), style='List Bullet')

    def shade(cell, fill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), fill)
        tcPr.append(shd)

    def table(headers, rows, widths_cm, empty_text):
        if not rows:
            para(empty_text, italic=True, color=GREY)
            return
        tb = doc.add_table(rows=1, cols=len(headers))
        tb.style = 'Table Grid'
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        tb.autofit = False
        hdr = tb.rows[0]
        trPr = hdr._tr.get_or_add_trPr()
        th = OxmlElement('w:tblHeader')
        th.set(qn('w:val'), 'true')
        trPr.append(th)
        for i, h in enumerate(headers):
            c = hdr.cells[i]
            c.width = Cm(widths_cm[i])
            c.text = ''
            run = c.paragraphs[0].add_run(h)
            run.bold = True
            run.font.size = Pt(9.5)
            shade(c, 'E8EDF5')
        for row in rows:
            cells = tb.add_row().cells
            for i, val in enumerate(row):
                cells[i].width = Cm(widths_cm[i])
                cells[i].text = ''
                run = cells[i].paragraphs[0].add_run(str(val or ''))
                run.font.size = Pt(9.5)
        for i, w in enumerate(widths_cm):
            tb.columns[i].width = Cm(w)

    decisions = [d for d in (clean_decision(x) for x in (rep.get('decisions') or [])) if d]
    tasks = [a for a in (clean_action(x) for x in (rep.get('action_items') or [])) if a]
    qrows = question_rows(rep)
    slides = rep.get('slides') or []

    # ---- title block
    doc.add_paragraph(rep.get('title') or 'Meeting', style='Title')
    para(f"SessionMind AI Meeting Report · Project: {rep.get('project', '')}", color=(0x55, 0x55, 0x55))
    para(f"Started: {nice_dt(rep.get('started_at'))}   Ended: {nice_dt(rep.get('ended_at')) if rep.get('ended_at') else '—'}",
         size=9, color=(0x55, 0x55, 0x55))
    for src in rep.get('sources') or []:
        para(f'Source: {src}', size=9, color=(0x55, 0x55, 0x55))

    doc.add_heading('Executive Summary', level=1)
    para(rep.get('executive_summary') or 'No summary generated.')
    doc.add_heading('Meeting Overview', level=1)
    para(rep.get('meeting_overview') or '')
    doc.add_heading('Participants', level=1)
    bullets(rep.get('participants') or [], empty='Unknown (not stated in transcript).')
    doc.add_heading('Key Discussion Points', level=1)
    bullets(rep.get('key_discussion_points') or [])

    doc.add_heading('Decisions', level=1)
    if decisions:
        for d in decisions:
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(d['decision']).bold = True
            meta = '  '.join(x for x in (f"[{d['timestamp']}]" if d['timestamp'] else '', f"— {d['evidence']}" if d['evidence'] else '') if x)
            if meta:
                ev = p.add_run('  ' + meta)
                ev.italic = True
                ev.font.size = Pt(9)
    else:
        para('No decisions were recorded.', italic=True, color=GREY)

    doc.add_heading('Action Items', level=1)
    table(['Owner', 'Task', 'Deadline', 'Status', 'When'],
          [[a['owner'], a['task'], a['deadline'], a['status'], a['timestamp']] for a in tasks],
          [2.8, 7.2, 2.5, 2.2, 2.3], 'No action items were recorded.')

    doc.add_heading('Open Questions', level=1)
    if qrows:
        counts = {}
        for q in qrows:
            counts[q['status']] = counts.get(q['status'], 0) + 1
        summary = ' · '.join(f'{v} {k.lower()}' for k, v in sorted(counts.items()))
        para(f'{len(qrows)} question(s) recorded — {summary}.', size=9.5, color=(0x55, 0x55, 0x55))
    table(['#', 'Question', 'Raised', 'Status', 'Evidence'],
          [[i + 1, q['question'], q['raised_by'] or q['timestamp'], q['status'], q['evidence']] for i, q in enumerate(qrows)],
          [0.9, 6.0, 2.4, 2.6, 5.1], 'No open questions were recorded.')

    doc.add_heading('Risks / Blockers', level=1)
    bullets([f"{r.get('risk')} — {r.get('evidence', '')}" if isinstance(r, dict) else str(r)
             for r in (rep.get('risks') or []) if r])

    doc.add_heading('Mind Map', level=1)
    if mind_view(tree) == 'tree':
        intro = para('Tree view: root = meeting, branches = themes, leaves = points, decisions, tasks and questions.',
                     size=9, color=(0x55, 0x55, 0x55))
        intro.paragraph_format.keep_with_next = True
        lines = tree_lines(tree)
        for i, line in enumerate(lines):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = i < len(lines) - 1
            run = p.add_run(line)
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.bold = (i == 0)
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
    else:
        pts = key_points(tree)
        if not pts:
            para('No key points were recorded.', italic=True, color=GREY)
        for theme, kids in pts:
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(theme).bold = True
            for k in kids:
                doc.add_paragraph(k, style='List Bullet 2')

    doc.add_heading('Important Moments', level=1)
    moments = [m for m in (rep.get('important_moments') or []) if isinstance(m, dict)]
    bullets([f"{m.get('timestamp', '')} — {m.get('description', '')}"
             + (f" [{m['evidence']}]" if m.get('evidence') else '') for m in moments])

    doc.add_heading('Changes From Previous Meeting', level=1)
    bullets(rep.get('changes_from_previous_meeting') or [], empty='No previous meeting available.')

    if slides:
        doc.add_heading('Presentation Slides', level=1)
        for s in slides:
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.add_run(f"Slide {s.get('slide_no', '')}: {s.get('title') or '(untitled)'}").bold = True
            if s.get('text'):
                para(s['text'][:1500], size=9.5)
            if s.get('notes'):
                para('Speaker notes: ' + s['notes'][:800], italic=True, size=9, color=GREY)

    doc.add_heading('Meeting Screenshots', level=1)
    shown = 0
    for s in screenshots[:12]:
        p = Path(base_dir) / s['path']
        if not p.exists():
            continue
        try:
            doc.add_picture(str(p), width=Cm(15))
            para(f"Captured at {clock(s['timestamp'])}", size=8, color=(0x77, 0x77, 0x77))
            shown += 1
        except Exception:
            pass
    if not shown:
        para('No screenshots captured.', italic=True, color=GREY)

    doc.add_heading('Evidence Transcript', level=1)
    for x in transcript:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        ts = p.add_run(f"[{clock(x['timestamp'])}] {x['speaker']}  ")
        ts.bold = True
        ts.font.size = Pt(9)
        ts.font.color.rgb = RGBColor(*GREY)
        p.add_run(x['text']).font.size = Pt(10)

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


# ---------------------------------------------------------------- HTML

def _tree_html(node):
    kids = node.get('children') or []
    inner = f'<ul>{"".join(_tree_html(k) for k in kids)}</ul>' if kids else ''
    return f'<li><span>{H(node["label"])}</span>{inner}</li>'


def mind_html(tree):
    if mind_view(tree) == 'tree':
        return f'<div class="tree"><ul>{_tree_html(tree)}</ul></div>'
    pts = key_points(tree)
    if not pts:
        return '<p>No key points were recorded.</p>'
    items = ''.join(
        f'<li><b>{H(theme)}</b>' + (f'<ul>{"".join("<li>" + H(k) + "</li>" for k in kids)}</ul>' if kids else '') + '</li>'
        for theme, kids in pts)
    return f'<ul class="points">{items}</ul>'


HTML_CSS = (
    'body{font:15px system-ui;max-width:1000px;margin:40px auto;color:#18202a;line-height:1.55;padding:0 16px}'
    'h1{margin-bottom:4px}h2{margin-top:30px;border-bottom:1px solid #ddd;padding-bottom:7px}'
    'table{width:100%;border-collapse:collapse}th,td{border:1px solid #ddd;padding:8px;text-align:left;vertical-align:top}'
    'th{background:#eef2f8}'
    'figure{display:inline-block;width:46%;vertical-align:top}img{max-width:100%;border:1px solid #ddd}'
    '.tree ul{list-style:none;margin:0;padding-left:22px;border-left:2px solid #c9d2e0}'
    '.tree>ul{border-left:none;padding-left:0}'
    '.tree li{position:relative;padding:4px 0 4px 14px}'
    '.tree li::before{content:"";position:absolute;left:0;top:14px;width:12px;border-top:2px solid #c9d2e0}'
    '.tree>ul>li::before{display:none}'
    '.tree span{display:inline-block;border:1px solid #aab6c8;border-radius:8px;padding:3px 10px;background:#f7f9fc}'
    '.tree>ul>li>span{font-weight:700;background:#e8eefc}'
    '.slide{border-left:3px solid #c9d2e0;padding-left:12px;margin:14px 0}'
)


def render_html(rep, transcript, screenshots, path, tree=None):
    tree = tree or build_mind_map(rep)
    tasks = [a for a in (clean_action(x) for x in (rep.get('action_items') or [])) if a]
    decisions = [d for d in (clean_decision(x) for x in (rep.get('decisions') or [])) if d]
    qrows = question_rows(rep)
    slides = rep.get('slides') or []

    task_rows = ''.join(
        f'<tr><td>{H(x["owner"])}</td><td>{H(x["task"])}</td><td>{H(x["deadline"])}</td><td>{H(x["status"])}</td><td>{H(x["timestamp"])}</td></tr>'
        for x in tasks) or '<tr><td colspan="5">No action items were recorded.</td></tr>'
    q_rows = ''.join(
        f'<tr><td>{i + 1}</td><td>{H(q["question"])}</td><td>{H(q["raised_by"] or q["timestamp"])}</td><td>{H(q["status"])}</td><td>{H(q["evidence"])}</td></tr>'
        for i, q in enumerate(qrows)) or '<tr><td colspan="5">No open questions were recorded.</td></tr>'
    dec_html = ''.join(
        f'<li><b>{H(d["decision"])}</b> {H("[" + d["timestamp"] + "]") if d["timestamp"] else ""} {H(d["evidence"])}</li>'
        for d in decisions) or '<li>No decisions were recorded.</li>'
    shots = ''.join(
        f'<figure><img src="/media/{H(Path(x["path"]).name)}"><figcaption>{clock(x["timestamp"])}</figcaption></figure>'
        for x in screenshots)

    def bullets(a, key=None):
        return ''.join('<li>' + H(x if key is None else x.get(key, '')) + '</li>' for x in a) or '<li>None recorded.</li>'

    slides_html = ''.join(
        f'<div class="slide"><b>Slide {H(s.get("slide_no"))}: {H(s.get("title") or "(untitled)")}</b>'
        + (f'<p>{H(s["text"][:1500])}</p>' if s.get('text') else '')
        + (f'<p><i>Speaker notes: {H(s["notes"][:800])}</i></p>' if s.get('notes') else '') + '</div>'
        for s in slides)
    transcript_html = ''.join(
        f'<p><b>[{clock(x["timestamp"])}] {H(x["speaker"])}</b> — {H(x["text"])}</p>' for x in transcript)
    risks = ''.join(f'<li>{H(x.get("risk"))} — {H(x.get("evidence"))}</li>'
                    for x in rep.get('risks', []) if isinstance(x, dict)) or '<li>None recorded.</li>'
    moments = ''.join(f'<li><b>{H(x.get("timestamp"))}</b> — {H(x.get("description"))} [{H(x.get("evidence"))}]</li>'
                      for x in rep.get('important_moments', []) if isinstance(x, dict)) or '<li>None recorded.</li>'
    sources = ''.join(f'<li>{H(s)}</li>' for s in rep.get('sources') or [])
    slides_section = ('<h2>Presentation Slides</h2>' + slides_html) if slides else ''

    body = (f'<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
            f'<title>{H(rep.get("title"))} - SessionMind</title><style>{HTML_CSS}</style></head><body>'
            f'<h1>{H(rep.get("title"))}</h1><div>SessionMind AI Meeting Report · {H(rep.get("project"))}</div><ul>{sources}</ul>'
            f'<h2>Executive Summary</h2><p>{H(rep.get("executive_summary"))}</p>'
            f'<h2>Meeting Overview</h2><p>{H(rep.get("meeting_overview"))}</p>'
            f'<h2>Participants</h2><ul>{bullets(rep.get("participants", []))}</ul>'
            f'<h2>Key Discussion Points</h2><ul>{bullets(rep.get("key_discussion_points", []))}</ul>'
            f'<h2>Decisions</h2><ul>{dec_html}</ul>'
            f'<h2>Action Items</h2><table><tr><th>Owner</th><th>Task</th><th>Deadline</th><th>Status</th><th>When</th></tr>{task_rows}</table>'
            f'<h2>Open Questions ({len(qrows)})</h2><table><tr><th>#</th><th>Question</th><th>Raised</th><th>Status</th><th>Evidence</th></tr>{q_rows}</table>'
            f'<h2>Risks / Blockers</h2><ul>{risks}</ul>'
            f'<h2>Mind Map</h2>{mind_html(tree)}'
            f'<h2>Important Moments</h2><ul>{moments}</ul>'
            f'<h2>Changes From Previous Meeting</h2><ul>{bullets(rep.get("changes_from_previous_meeting", []))}</ul>'
            f'{slides_section}'
            f'<h2>Meeting Screenshots</h2>{shots or "<p>No screenshots captured.</p>"}'
            f'<h2>Evidence Transcript</h2>{transcript_html}</body></html>')
    Path(path).write_text(body, encoding='utf-8')
