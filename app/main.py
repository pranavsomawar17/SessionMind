from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from pathlib import Path
from datetime import datetime, timezone
import sqlite3, os, json, time, subprocess, shutil, re, html, difflib
import requests
from .report_render import (build_mind_map, clean_action, clean_decision, clean_question, clock, merge_items, mind_view, nice_dt, render_docx, render_html)

BASE=Path(__file__).resolve().parents[1]
try:
    from dotenv import load_dotenv
    load_dotenv(BASE/'.env')
except Exception:
    pass
DATA=BASE/'data'; MEDIA=DATA/'media'; DATA.mkdir(exist_ok=True); MEDIA.mkdir(exist_ok=True)
DB=DATA/'sessionmind.db'
app=FastAPI(title='SessionMind',version='3.0.0')
app.mount('/static',StaticFiles(directory=str(BASE/'static')),name='static')
app.mount('/media',StaticFiles(directory=str(MEDIA)),name='media')

class ProjectIn(BaseModel): name:str
class MeetingIn(BaseModel): project_id:int; title:str='Meeting'


def db():
    c=sqlite3.connect(DB); c.row_factory=sqlite3.Row
    c.execute('PRAGMA busy_timeout=5000')
    c.executescript('''
    CREATE TABLE IF NOT EXISTS projects(id INTEGER PRIMARY KEY AUTOINCREMENT,name TEXT NOT NULL,created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS meetings(id INTEGER PRIMARY KEY AUTOINCREMENT,project_id INTEGER NOT NULL,title TEXT NOT NULL,status TEXT NOT NULL,started_at TEXT NOT NULL,ended_at TEXT,report TEXT);
    CREATE TABLE IF NOT EXISTS transcript(id INTEGER PRIMARY KEY AUTOINCREMENT,meeting_id INTEGER NOT NULL,speaker TEXT NOT NULL,text TEXT NOT NULL,timestamp REAL NOT NULL,source TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS screenshots(id INTEGER PRIMARY KEY AUTOINCREMENT,meeting_id INTEGER NOT NULL,path TEXT NOT NULL,timestamp REAL NOT NULL);
    CREATE TABLE IF NOT EXISTS events(id INTEGER PRIMARY KEY AUTOINCREMENT,meeting_id INTEGER NOT NULL,type TEXT NOT NULL,content TEXT NOT NULL,assignee TEXT,deadline TEXT,timestamp REAL NOT NULL,confidence REAL NOT NULL,status TEXT NOT NULL,evidence TEXT DEFAULT '');
    CREATE TABLE IF NOT EXISTS reports(id INTEGER PRIMARY KEY AUTOINCREMENT,meeting_id INTEGER UNIQUE NOT NULL,html_path TEXT,pdf_path TEXT,created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS slides(id INTEGER PRIMARY KEY AUTOINCREMENT,meeting_id INTEGER NOT NULL,slide_no INTEGER NOT NULL,title TEXT,text TEXT,notes TEXT);
    CREATE TABLE IF NOT EXISTS materials(id INTEGER PRIMARY KEY AUTOINCREMENT,meeting_id INTEGER NOT NULL,kind TEXT NOT NULL,filename TEXT,detail TEXT,created_at TEXT NOT NULL);
    ''')
    # Migration-safe additions for meetings created before live-insights support existed.
    for table,col,decl in (('meetings','last_analyzed_ts','REAL DEFAULT 0'),('meetings','current_topic','TEXT'),('reports','docx_path','TEXT')):
        try: c.execute(f'ALTER TABLE {table} ADD COLUMN {col} {decl}')
        except sqlite3.OperationalError: pass
    return c

def now(): return datetime.now(timezone.utc).isoformat()
def env_keys(name): return [x.strip() for x in os.getenv(name,'').split(',') if x.strip()]
def groq_keys(): return env_keys('GROQ_API_KEY') or env_keys('GROQ_API_KEYS')

def require_config():
    missing=[]
    if not groq_keys(): missing.append('GROQ_API_KEY')
    if not shutil.which('ffmpeg'): missing.append('ffmpeg')
    return missing

def groq_transcribe(wav:Path):
    last=''
    for key in groq_keys():
        try:
            with wav.open('rb') as f:
                r=requests.post('https://api.groq.com/openai/v1/audio/transcriptions',headers={'Authorization':f'Bearer {key}'},files={'file':(wav.name,f,'audio/wav')},data={'model':os.getenv('GROQ_TRANSCRIPTION_MODEL','whisper-large-v3-turbo'),'response_format':'verbose_json','temperature':'0'},timeout=180)
            if r.status_code in (401,403): last=f'Groq authentication failed: {r.text[:500]}'; continue
            if r.status_code==429: last=f'Groq quota/rate limit: {r.text[:500]}'; continue
            r.raise_for_status(); return r.json()
        except Exception as e: last=str(e)
    raise RuntimeError(last or 'Groq transcription failed')

def extract_wav(src:Path)->Path:
    ff=shutil.which('ffmpeg')
    if not ff: raise RuntimeError('FFmpeg is not installed. Fedora: sudo dnf install ffmpeg')
    out=src.with_suffix('.wav')
    # Robustly decode browser WebM/Opus. Explicitly ignore video because the upload is audio-only.
    cmd=[ff,'-hide_banner','-loglevel','error','-y','-i',str(src),'-map','0:a:0','-vn','-ac','1','-ar','16000','-c:a','pcm_s16le',str(out)]
    p=subprocess.run(cmd,capture_output=True,text=True,timeout=180)
    if p.returncode or not out.exists() or out.stat().st_size<1000:
        raise RuntimeError('FFmpeg could not decode this audio recording. Browser capture may have produced an empty stream. '+p.stderr[-800:])
    return out

def groq_chat_model(key):
    """Pick a strong currently available Groq chat model.
    A configured GROQ_CHAT_MODEL wins when it is available; otherwise use a
    sensible preference order and finally any available text-generation model.
    """
    configured=os.getenv('GROQ_CHAT_MODEL','llama-3.3-70b-versatile').strip()
    try:
        r=requests.get('https://api.groq.com/openai/v1/models',
                       headers={'Authorization':f'Bearer {key}','accept':'application/json'},
                       timeout=30)
        if r.ok:
            ids=[x.get('id','') for x in r.json().get('data',[]) if x.get('id')]
            if configured in ids:
                return configured
            preferred=[
                'llama-3.3-70b-versatile',
                'openai/gpt-oss-120b',
                'qwen/qwen3-32b',
                'llama-3.1-70b-versatile',
                'llama-3.1-8b-instant'
            ]
            for model in preferred:
                if model in ids:
                    return model
            # Avoid embedding/whisper/guard models when possible.
            blocked=('whisper','embed','guard','safeguard','compound')
            text_models=[x for x in ids if not any(b in x.lower() for b in blocked)]
            if text_models:
                return text_models[0]
    except Exception:
        pass
    return configured

def groq_json(system,user):
    """Run structured meeting analysis through Groq's OpenAI-compatible API."""
    last=''
    for key in groq_keys():
        model=groq_chat_model(key)
        payload={
            'model':model,
            'messages':[{'role':'system','content':system},{'role':'user','content':user}],
            'temperature':0.15,
            'max_tokens':8000,
            'response_format':{'type':'json_object'}
        }
        try:
            r=requests.post(
                'https://api.groq.com/openai/v1/chat/completions',
                headers={'Authorization':f'Bearer {key}','Content-Type':'application/json'},
                json=payload,timeout=300
            )
            if r.status_code in (401,403):
                last=f'Groq authentication failed: {r.text[:1200]}'; continue
            if r.status_code==429:
                last=f'Groq rate/quota limit: {r.text[:1200]}'; continue
            if not r.ok:
                try: detail=r.json().get('error',{}).get('message') or r.text
                except Exception: detail=r.text
                last=f'Groq API error {r.status_code} using model {model}: {detail[:1500]}'
                continue
            data=r.json()
            text=data.get('choices',[{}])[0].get('message',{}).get('content','').strip()
            try:
                return json.loads(text)
            except Exception:
                m=re.search(r'\{.*\}',text,re.S)
                if m:
                    return json.loads(m.group(0))
                raise RuntimeError('Groq returned non-JSON output')
        except Exception as e:
            last=str(e)
    raise RuntimeError(last or 'Groq meeting analysis failed')

def groq_vision_model(key):
    """Best-effort pick of a currently available Groq vision-capable model.
    Returns None if the account has no vision model available — callers must
    treat that as 'skip visual analysis', never as a fatal error."""
    try:
        r=requests.get('https://api.groq.com/openai/v1/models',
                       headers={'Authorization':f'Bearer {key}','accept':'application/json'},
                       timeout=30)
        if r.ok:
            ids=[x.get('id','') for x in r.json().get('data',[]) if x.get('id')]
            preferred=[
                'meta-llama/llama-4-scout-17b-16e-instruct',
                'meta-llama/llama-4-maverick-17b-128e-instruct',
                'llama-3.2-90b-vision-preview',
                'llama-3.2-11b-vision-preview',
            ]
            for model in preferred:
                if model in ids: return model
            for i in ids:
                low=i.lower()
                if 'vision' in low or 'scout' in low or 'maverick' in low:
                    return i
    except Exception:
        pass
    return None

def groq_vision_read(image_bytes):
    """Read live captions or slide/document text out of a screenshot.
    Best-effort: returns None on any failure or when no vision model is
    available, so a screenshot never blocks or breaks the meeting."""
    import base64
    b64=base64.b64encode(image_bytes).decode()
    prompt=('This image is a screenshot of a live video-call screen (e.g. Zoom or Google Meet). '
            'If live captions/subtitles are visible, transcribe them exactly, including the speaker '
            'name if one is labeled next to the caption. Otherwise, if a presentation slide or shared '
            'document with meaningful text is visible, extract its key heading/text briefly. '
            'Otherwise reply with exactly: NONE. Reply with only the transcription or extracted text, '
            'nothing else — no preamble, no quotes.')
    for key in groq_keys():
        model=groq_vision_model(key)
        if not model: continue
        try:
            payload={
                'model':model,
                'messages':[{'role':'user','content':[
                    {'type':'text','text':prompt},
                    {'type':'image_url','image_url':{'url':f'data:image/jpeg;base64,{b64}'}}
                ]}],
                'temperature':0,'max_tokens':300
            }
            r=requests.post('https://api.groq.com/openai/v1/chat/completions',
                            headers={'Authorization':f'Bearer {key}','Content-Type':'application/json'},
                            json=payload,timeout=60)
            if not r.ok: continue
            text=r.json().get('choices',[{}])[0].get('message',{}).get('content','').strip()
            if text and text.upper()!='NONE':
                return text
            return None
        except Exception:
            continue
    return None

def save_transcript(mid,text,ts,speaker='Meeting Speaker',source='groq_whisper'):
    text=(text or '').strip()
    if not text:return []
    c=db(); cur=c.execute('INSERT INTO transcript(meeting_id,speaker,text,timestamp,source) VALUES(?,?,?,?,?)',(mid,speaker,text,ts,source)); c.commit(); row=c.execute('SELECT * FROM transcript WHERE id=?',(cur.lastrowid,)).fetchone(); c.close(); return dict(row)

def parse_ts(v, fallback=0):
    try:return float(v)
    except:return fallback

def previous_context(mid):
    c=db(); m=c.execute('SELECT project_id FROM meetings WHERE id=?',(mid,)).fetchone()
    if not m:c.close();return ''
    rows=c.execute('SELECT title,report FROM meetings WHERE project_id=? AND id<? AND status="ended" ORDER BY id DESC LIMIT 3',(m['project_id'],mid)).fetchall(); c.close()
    out=[]
    for r in rows:
        if not r['report']:continue
        try:
            x=json.loads(r['report']); out.append({'title':r['title'],'summary':x.get('executive_summary',''),'decisions':x.get('decisions',[]),'action_items':x.get('action_items',[]),'open_questions':x.get('open_questions',[])})
        except:pass
    return json.dumps(out,ensure_ascii=False)[:16000]

@app.on_event('startup')
def startup():
    db().close()

@app.get('/')
def root(): return FileResponse(BASE/'static/index.html')

@app.get('/api/status')
def status():
    missing=require_config()
    return {'ok':not missing,'missing':missing,'groq_configured':bool(groq_keys()),'llm_configured':bool(groq_keys()),'ffmpeg':bool(shutil.which('ffmpeg')),'capture':'browser screen + system audio','transcription':'Groq Whisper','analysis':'Groq LLM'}

@app.get('/api/projects')
def projects():
    c=db(); rows=c.execute('SELECT * FROM projects ORDER BY id DESC').fetchall(); c.close(); return [dict(x) for x in rows]
@app.post('/api/projects')
def create_project(x:ProjectIn):
    name=x.name.strip() or 'My Project'; c=db(); cur=c.execute('INSERT INTO projects(name,created_at) VALUES(?,?)',(name,now())); c.commit(); row=c.execute('SELECT * FROM projects WHERE id=?',(cur.lastrowid,)).fetchone(); c.close(); return dict(row)
@app.get('/api/projects/{pid}/meetings')
def meetings(pid:int):
    c=db(); rows=c.execute('SELECT id,title,status,started_at,ended_at FROM meetings WHERE project_id=? ORDER BY id DESC',(pid,)).fetchall(); c.close(); return [dict(x) for x in rows]
@app.post('/api/meetings')
def create_meeting(x:MeetingIn):
    c=db(); p=c.execute('SELECT id FROM projects WHERE id=?',(x.project_id,)).fetchone()
    if not p:c.close(); raise HTTPException(404,'Project not found')
    cur=c.execute('INSERT INTO meetings(project_id,title,status,started_at) VALUES(?,?,?,?)',(x.project_id,x.title.strip() or 'Meeting','live',now())); c.commit(); row=c.execute('SELECT * FROM meetings WHERE id=?',(cur.lastrowid,)).fetchone(); c.close(); return dict(row)

@app.post('/api/meetings/{mid}/audio')
def audio(mid:int,audio:UploadFile=File(...),started_at:float|None=Form(None),speaker:str|None=Form(None)):
    # Plain `def`, not `async def`: FastAPI runs sync endpoints in a worker
    # thread automatically, so the ffmpeg decode + Groq call below (which can
    # take several seconds) no longer blocks the event loop for every other
    # in-flight request (screenshots, status checks, other audio chunks).
    c=db(); m=c.execute('SELECT status,started_at FROM meetings WHERE id=?',(mid,)).fetchone(); c.close()
    if not m: raise HTTPException(404,'Meeting not found')
    if m['status']!='live': raise HTTPException(409,'Meeting is not live')
    data=audio.file.read()
    if len(data)<4000: return {'ok':True,'ignored':True,'reason':'small audio'}
    ext=Path(audio.filename or 'capture.webm').suffix.lower() or '.webm'
    src=MEDIA/f'{mid}_{int(time.time()*1000)}{ext}'; src.write_bytes(data)
    wav=None
    try:
        wav=extract_wav(src); result=groq_transcribe(wav)
        base=started_at or time.time()-float(result.get('duration') or 0)
        created=[]
        segs=result.get('segments') or []
        label=(speaker or '').strip() or 'Meeting Speaker'
        if segs:
            for s in segs:
                txt=(s.get('text') or '').strip()
                if txt: created.append(save_transcript(mid,txt,base+parse_ts(s.get('start'),0),label))
        elif result.get('text'): created.append(save_transcript(mid,result['text'],base,label))
        return {'ok':True,'text':result.get('text',''),'segments':created}
    except Exception as e:
        raise HTTPException(422,f'Audio processing error: {e}')
    finally:
        for p in (src,wav):
            if p:
                try:p.unlink()
                except:pass

@app.post('/api/meetings/{mid}/screenshot')
def screenshot(mid:int,image:UploadFile=File(...),timestamp:float|None=Form(None)):
    c=db(); ok=c.execute('SELECT id FROM meetings WHERE id=?',(mid,)).fetchone(); c.close()
    if not ok:raise HTTPException(404,'Meeting not found')
    ts=timestamp or time.time()
    data=image.file.read()
    p=MEDIA/f'{mid}_{int(ts*1000)}.jpg'; p.write_bytes(data)
    c=db(); cur=c.execute('INSERT INTO screenshots(meeting_id,path,timestamp) VALUES(?,?,?)',(mid,str(p.relative_to(BASE)),ts)); c.commit(); row=c.execute('SELECT * FROM screenshots WHERE id=?',(cur.lastrowid,)).fetchone(); c.close()
    # Best-effort: read captions/slide text out of the frame and fold it into
    # the transcript as evidence. Never lets a vision failure break capture.
    try:
        text=groq_vision_read(data)
        if text: save_transcript(mid,text,ts,'Screen',source='groq_vision')
    except Exception:
        pass
    return dict(row)

@app.post('/api/meetings/{mid}/analyze')
def analyze(mid:int):
    """Incremental live extraction: pulls transcript lines added since the
    last call, asks Groq for any new tasks/decisions/questions/topic change,
    and stores them in the existing (previously unused) events table."""
    c=db()
    m=c.execute('SELECT * FROM meetings WHERE id=?',(mid,)).fetchone()
    if not m: c.close(); raise HTTPException(404,'Meeting not found')
    rows=c.execute('SELECT * FROM transcript WHERE meeting_id=? AND timestamp>? ORDER BY timestamp',(mid,m['last_analyzed_ts'] or 0)).fetchall()
    if not rows:
        events=[dict(x) for x in c.execute('SELECT * FROM events WHERE meeting_id=? ORDER BY id',(mid,)).fetchall()]
        c.close()
        return {'events':events,'current_topic':m['current_topic'],'new':0}
    window='\n'.join(f'[{x["timestamp"]:.1f}] {x["speaker"]}: {x["text"]}' for x in rows)
    existing=c.execute('SELECT type,content FROM events WHERE meeting_id=?',(mid,)).fetchall()
    existing_desc='\n'.join(f'- ({x["type"]}) {x["content"]}' for x in existing) or 'None yet'
    system=('You are the real-time extraction engine for SessionMind. Read the new transcript segment '
            'and extract structured events. Extract a "decision" only when there is clear agreement or '
            'confirmation, not a mere suggestion. Extract a "task" only when someone is clearly assigned '
            'to or commits to an action. Extract a "question" for every question or open issue raised in the '
            'segment that is not answered within that same segment, including questions the group defers, '
            'items marked "we need to check/confirm", and unknown owners or deadlines that were flagged. '
            'Do not invent information not present in the transcript. Do not repeat already-detected '
            'events. Return ONLY valid JSON with this exact shape: '
            '{"topic":string|null,"events":[{"type":"task|decision|question","description":"","assignee":string|null,"deadline":string|null,"confidence":0.0}]}')
    user=f'CURRENT TOPIC (may be unchanged): {m["current_topic"] or "unknown"}\n\nNEW TRANSCRIPT SEGMENT:\n{window}\n\nALREADY-DETECTED EVENTS (do not repeat):\n{existing_desc}'
    try:
        result=groq_json(system,user)
    except Exception as e:
        c.close(); raise HTTPException(502,f'Live analysis failed: {e}')
    last_ts=rows[-1]['timestamp']
    topic=result.get('topic') or m['current_topic']
    c.execute('UPDATE meetings SET last_analyzed_ts=?,current_topic=? WHERE id=?',(last_ts,topic,mid))
    for ev in (result.get('events') or []):
        content=(ev.get('description') or '').strip()
        etype=ev.get('type') if ev.get('type') in ('task','decision','question') else None
        if not content or not etype: continue
        dup=c.execute('SELECT id FROM events WHERE meeting_id=? AND type=? AND lower(content)=lower(?)',(mid,etype,content)).fetchone()
        if dup: continue
        c.execute('INSERT INTO events(meeting_id,type,content,assignee,deadline,timestamp,confidence,status,evidence) VALUES(?,?,?,?,?,?,?,?,?)',
                  (mid,etype,content,ev.get('assignee'),ev.get('deadline'),last_ts,float(ev.get('confidence') or 0.7),'detected',window[-500:]))
    c.commit()
    events=[dict(x) for x in c.execute('SELECT * FROM events WHERE meeting_id=? ORDER BY id',(mid,)).fetchall()]
    c.close()
    return {'events':events,'current_topic':topic,'new':len(rows)}

@app.get('/api/meetings/{mid}/events')
def list_events(mid:int):
    c=db(); ok=c.execute('SELECT id FROM meetings WHERE id=?',(mid,)).fetchone()
    if not ok: c.close(); raise HTTPException(404,'Meeting not found')
    rows=c.execute('SELECT * FROM events WHERE meeting_id=? ORDER BY id',(mid,)).fetchall(); c.close(); return [dict(x) for x in rows]

class EventStatus(BaseModel): status:str
@app.patch('/api/meetings/{mid}/events/{eid}')
def update_event(mid:int,eid:int,x:EventStatus):
    if x.status not in ('detected','confirmed','dismissed'): raise HTTPException(400,'Invalid status')
    c=db(); cur=c.execute('UPDATE events SET status=? WHERE id=? AND meeting_id=?',(x.status,eid,mid))
    if cur.rowcount==0: c.close(); raise HTTPException(404,'Event not found')
    c.commit(); row=c.execute('SELECT * FROM events WHERE id=?',(eid,)).fetchone(); c.close(); return dict(row)

@app.post('/api/meetings/{mid}/end')
def end_meeting(mid:int):
    c=db(); m=c.execute('SELECT * FROM meetings WHERE id=?',(mid,)).fetchone(); c.close()
    if not m:raise HTTPException(404,'Meeting not found')
    c=db(); c.execute('UPDATE meetings SET status="ended",ended_at=? WHERE id=?',(now(),mid)); c.commit(); c.close()
    try:return build_report(mid)
    except Exception as e:
        # Keep the meeting ended and expose a useful error instead of an opaque 500.
        raise HTTPException(502,f'Report generation failed: {e}')

SCAN_SYSTEM=('You extract the important items from ONE segment of a meeting transcript. Be thorough: a missed item is worse than a borderline one, '
 'but never invent anything that is not in the segment. '
 'DECISION: any point the group agreed or settled, e.g. "agreed", "let us go with X", "yes, do that", "we will use X". Not mere suggestions. '
 'ACTION ITEM: any concrete task someone was asked to do or volunteered for, e.g. "can you check", "I will send", "we need to update X". '
 'owner = the named person or "Unknown"; deadline = the stated date or "Unknown". '
 'OPEN QUESTION: any question raised that the segment does not fully answer, any item deferred ("later", "let us check", "we need to confirm"), '
 'any undecided option, or missing information. '
 'Use empty lists when a category has nothing. Return ONLY valid JSON: '
 '{"decisions":[{"decision":"","evidence":"","timestamp":""}],'
 '"action_items":[{"owner":"","task":"","deadline":"","evidence":"","timestamp":""}],'
 '"open_questions":[{"question":"","raised_by":"","timestamp":"","status":"Unanswered|Partially answered|Deferred","evidence":""}]}. '
 'evidence quotes or closely paraphrases the transcript line; timestamp is that line\'s [mm:ss].')

def transcript_chunks(rows,limit=24000):
    """Yield the transcript in ~24k-char pieces, split on line boundaries."""
    chunk=[]; size=0
    for x in rows:
        line=f'[{clock(x["timestamp"])}] {x["speaker"]}: {x["text"]}'
        if chunk and size+len(line)>limit:
            yield '\n'.join(chunk); chunk=[]; size=0
        chunk.append(line); size+=len(line)+1
    if chunk: yield '\n'.join(chunk)

def scan_items(rows):
    """Full-transcript pass in chunks: returns (decisions, action_items, open_questions) candidates."""
    dec=[]; act=[]; qs=[]; failed=0; chunks=0
    for chunk in transcript_chunks(rows):
        chunks+=1
        try: res=groq_json(SCAN_SYSTEM,'TRANSCRIPT SEGMENT:\n'+chunk)
        except Exception as e:
            failed+=1; print(f'[SessionMind] scan chunk failed: {e}'); continue
        if not isinstance(res,dict): res={}
        dec+=[x for x in (clean_decision(i) for i in (res.get('decisions') or [])) if x]
        act+=[x for x in (clean_action(i) for i in (res.get('action_items') or [])) if x]
        qs+=[x for x in (clean_question(i) for i in (res.get('open_questions') or [])) if x]
    print(f'[SessionMind] scan: {chunks} chunk(s), {failed} failed; candidates dec={len(dec)} act={len(act)} q={len(qs)}')
    return merge_items('decision',dec), merge_items('task',act), merge_items('question',qs)

def budget_json(items,keys,limit):
    """JSON array of whole items, cut at a size budget (never mid-item)."""
    out=[]; used=0
    for it in items:
        s=json.dumps({k:it.get(k,'') for k in keys},ensure_ascii=False)
        if used+len(s)>limit: break
        out.append(s); used+=len(s)+2
    return '['+', '.join(out)+']'

def hydrate_report(raw):
    rep=json.loads(raw); tree=build_mind_map(rep)
    rep['mind_map_tree']=tree; rep['mind_view']=mind_view(tree); return rep

@app.get('/api/meetings/{mid}')
def get_meeting(mid:int):
    c=db(); m=c.execute('SELECT * FROM meetings WHERE id=?',(mid,)).fetchone()
    if not m:c.close();raise HTTPException(404,'Meeting not found')
    t=c.execute('SELECT * FROM transcript WHERE meeting_id=? ORDER BY timestamp',(mid,)).fetchall(); s=c.execute('SELECT * FROM screenshots WHERE meeting_id=? ORDER BY timestamp',(mid,)).fetchall(); c.close()
    report=None
    if m['report']:
        try: report=hydrate_report(m['report'])
        except Exception: report=None
    return {**dict(m),'transcript':[dict(x) for x in t],'screenshots':[dict(x) for x in s],'report_data':report}

# ---------------------------------------------------------------- optional materials: recording + slides

SEGMENT_SECONDS=480  # 8-minute audio pieces keep every Groq upload far below the file-size limit

def ffprobe_duration(p):
    try:
        out=subprocess.run(['ffprobe','-v','error','-show_entries','format=duration','-of','default=nw=1:nk=1',str(p)],capture_output=True,text=True,timeout=120).stdout.strip()
        return float(out)
    except Exception: return 0.0

def run_ff(args,timeout=3600):
    p=subprocess.run(['ffmpeg','-hide_banner','-loglevel','error','-y',*args],capture_output=True,text=True,timeout=timeout)
    if p.returncode: raise RuntimeError('FFmpeg failed: '+p.stderr[-600:])

def process_recording(mid,src:Path,replace_live:bool,t0:float):
    """Recording (video or audio) -> transcript lines (Groq Whisper, in 8-min pieces) + ~40 evenly spaced stills."""
    if not shutil.which('ffmpeg'): raise RuntimeError('FFmpeg is not installed. Fedora: sudo dnf install ffmpeg')
    work=MEDIA/f'rec_{mid}_{int(time.time()*1000)}'; work.mkdir(parents=True)
    try:
        duration=ffprobe_duration(src)
        audio=work/'audio.wav'
        run_ff(['-i',str(src),'-vn','-ac','1','-ar','16000','-c:a','pcm_s16le',str(audio)])
        if not audio.exists() or audio.stat().st_size<1000: raise RuntimeError('The recording has no readable audio track.')
        run_ff(['-i',str(audio),'-f','segment','-segment_time',str(SEGMENT_SECONDS),'-c','copy',str(work/'part_%03d.wav')])
        new_rows=[]
        for idx,part in enumerate(sorted(work.glob('part_*.wav'))):
            res=groq_transcribe(part)
            offset=idx*SEGMENT_SECONDS
            segs=res.get('segments') or []
            if segs:
                for sg in segs:
                    txt=(sg.get('text') or '').strip()
                    if txt: new_rows.append((txt,t0+offset+parse_ts(sg.get('start'),0)))
            elif (res.get('text') or '').strip():
                new_rows.append((res['text'].strip(),t0+offset))
        if not new_rows: raise RuntimeError('No speech was recognised in the recording.')
        # Stills for visual evidence (video only; audio-only files just skip this).
        frames=[]; interval=max(60.0,duration/40) if duration>0 else 60.0
        try:
            run_ff(['-i',str(src),'-vf',f'fps=1/{interval}','-frames:v','40','-q:v','4',str(work/'frame_%03d.jpg')],timeout=1800)
            frames=sorted(work.glob('frame_*.jpg'))
        except Exception as e: print(f'[SessionMind] no stills from recording: {e}')
        stamp=int(time.time()*1000)
        c=db()
        for i,f in enumerate(frames):
            dest=MEDIA/f'{mid}_rec_{stamp}_{i:03d}.jpg'; shutil.copy(f,dest)
            c.execute('INSERT INTO screenshots(meeting_id,path,timestamp) VALUES(?,?,?)',(mid,str(dest.relative_to(BASE)),t0+i*interval))
        if replace_live: c.execute("DELETE FROM transcript WHERE meeting_id=? AND source='groq_whisper'",(mid,))
        for txt,ts in new_rows:
            c.execute('INSERT INTO transcript(meeting_id,speaker,text,timestamp,source) VALUES(?,?,?,?,?)',(mid,'Recording',txt,ts,'recording'))
        c.commit(); c.close()
        return {'lines':len(new_rows),'minutes':round(duration/60,1),'stills':len(frames),'replaced_live':replace_live}
    finally:
        shutil.rmtree(work,ignore_errors=True)

def process_slides(mid,src:Path):
    """PowerPoint -> per-slide title, text, tables and speaker notes."""
    from pptx import Presentation
    try: prs=Presentation(str(src))
    except Exception as e: raise RuntimeError(f'Could not read this PowerPoint file: {e}')
    rows=[]
    for i,slide in enumerate(prs.slides,1):
        title=''
        try:
            if slide.shapes.title is not None and slide.shapes.title.has_text_frame: title=slide.shapes.title.text_frame.text.strip()
        except Exception: pass
        parts=[]
        for shp in slide.shapes:
            if getattr(shp,'has_text_frame',False) and shp.has_text_frame:
                t='\n'.join(p.text for p in shp.text_frame.paragraphs if p.text.strip()).strip()
                if t and t!=title: parts.append(t)
            if getattr(shp,'has_table',False) and shp.has_table:
                for row in shp.table.rows: parts.append(' | '.join(cell.text.strip() for cell in row.cells))
        notes=''
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes=slide.notes_slide.notes_text_frame.text.strip()
        rows.append((i,title,'\n'.join(parts),notes))
    c=db(); c.execute('DELETE FROM slides WHERE meeting_id=?',(mid,))
    for r in rows: c.execute('INSERT INTO slides(meeting_id,slide_no,title,text,notes) VALUES(?,?,?,?,?)',(mid,*r))
    c.commit(); c.close()
    return len(rows)

@app.post('/api/meetings/{mid}/materials')
def upload_material(mid:int,kind:str=Form(...),file:UploadFile=File(...),replace_live:str=Form('true')):
    if kind not in ('recording','slides'): raise HTTPException(400,'kind must be "recording" or "slides"')
    c=db(); m=c.execute('SELECT id,started_at FROM meetings WHERE id=?',(mid,)).fetchone(); c.close()
    if not m: raise HTTPException(404,'Meeting not found')
    name=Path(file.filename or 'upload').name
    suffix=Path(name).suffix.lower()
    if kind=='slides' and suffix!='.pptx':
        raise HTTPException(400,'Slides must be a .pptx file. In Google Slides or Keynote, export/download as PowerPoint first.')
    dest=MEDIA/f'{mid}_{kind}_{int(time.time()*1000)}{suffix or ".bin"}'
    with dest.open('wb') as f: shutil.copyfileobj(file.file,f)
    try:
        if kind=='recording':
            t0=datetime.fromisoformat(m['started_at']).timestamp()
            info=process_recording(mid,dest,str(replace_live).lower()!='false',t0)
            detail=f"{info['minutes']} min · {info['lines']} transcript lines · {info['stills']} stills"
        else:
            detail=f"{process_slides(mid,dest)} slides"
    except Exception as e:
        raise HTTPException(422,f'{kind.title()} processing failed: {e}')
    finally:
        dest.unlink(missing_ok=True)
    c=db(); c.execute('INSERT INTO materials(meeting_id,kind,filename,detail,created_at) VALUES(?,?,?,?,?)',(mid,kind,name,detail,now())); c.commit(); c.close()
    return {'ok':True,'kind':kind,'filename':name,'detail':detail}

# ---------------------------------------------------------------- report

def build_report(mid):
    c=db(); m=c.execute('SELECT * FROM meetings WHERE id=?',(mid,)).fetchone(); p=c.execute('SELECT * FROM projects WHERE id=?',(m['project_id'],)).fetchone(); t=c.execute('SELECT * FROM transcript WHERE meeting_id=? ORDER BY timestamp',(mid,)).fetchall(); s=c.execute('SELECT * FROM screenshots WHERE meeting_id=? ORDER BY timestamp',(mid,)).fetchall(); e=c.execute('SELECT * FROM events WHERE meeting_id=? ORDER BY timestamp',(mid,)).fetchall(); sl=c.execute('SELECT slide_no,title,text,notes FROM slides WHERE meeting_id=? ORDER BY slide_no',(mid,)).fetchall(); mats=c.execute('SELECT kind,filename,detail FROM materials WHERE meeting_id=? ORDER BY id',(mid,)).fetchall(); c.close()
    if not t: raise RuntimeError('No transcript was captured. Make sure the screen-share dialog has Share audio enabled, or upload a recording, and allow the meeting to run for at least 10 seconds.')
    # Timestamps are meeting-relative (mm:ss), not epoch seconds.
    try: t0=datetime.fromisoformat(m['started_at']).timestamp()
    except Exception: t0=t[0]['timestamp']
    rel=lambda ts: max(0.0,float(ts)-t0)
    rows=[{**dict(x),'timestamp':rel(x['timestamp'])} for x in t]
    shots=[{**dict(x),'timestamp':rel(x['timestamp'])} for x in s]
    slides=[dict(x) for x in sl]

    # Pass 1: full-transcript scan for decisions, action items and questions.
    scan_dec,scan_act,scan_q=scan_items(rows)
    # Live-detected items (from the Live insights panel) - dismissed ones are excluded.
    def live(kind):
        return [x for x in e if x['type']==kind and x['status']!='dismissed']
    live_dec=[clean_decision({'decision':x['content'],'evidence':x['evidence'] or '','timestamp':clock(rel(x['timestamp']))}) for x in live('decision')]
    live_act=[clean_action({'task':x['content'],'owner':x['assignee'],'deadline':x['deadline'],'evidence':x['evidence'] or '','timestamp':clock(rel(x['timestamp']))}) for x in live('task')]
    live_q=[clean_question({'question':x['content'],'evidence':x['evidence'] or '','timestamp':clock(rel(x['timestamp']))}) for x in live('question')]
    cand_dec=merge_items('decision',scan_dec,[x for x in live_dec if x])
    cand_act=merge_items('task',scan_act,[x for x in live_act if x])
    cand_q=merge_items('question',scan_q,[x for x in live_q if x])

    transcript='\n'.join(f'[{clock(x["timestamp"])}] {x["speaker"]}: {x["text"]}' for x in rows)[-80000:]
    slide_text='\n'.join(f'Slide {x["slide_no"]} ({x["title"] or "untitled"}): {x["text"]}' for x in slides)[:15000]
    cand_block=(f'CANDIDATE DECISIONS (JSON):\n{budget_json(cand_dec,("decision","evidence","timestamp"),7000)}\n\n'
                f'CANDIDATE ACTION ITEMS (JSON):\n{budget_json(cand_act,("owner","task","deadline","evidence","timestamp"),7000)}\n\n'
                f'CANDIDATE OPEN QUESTIONS (JSON):\n{budget_json(cand_q,("question","status","raised_by","timestamp","evidence"),7000)}')
    system=('You are SessionMind, an evidence-first meeting analyst. Analyze the complete meeting transcript. Do not invent facts. '
            'Every extracted item must be traceable to transcript timestamps. Use "Unknown" when owner or deadline is not stated. '
            'DECISIONS: list points the group agreed or settled. Suggestions alone are not decisions. '
            'ACTION ITEMS: list every concrete task someone was asked to do or volunteered for, even if the owner or deadline is Unknown. '
            'OPEN QUESTIONS: list every unresolved question, deferred item, undecided option and missing information. '
            'You are given CANDIDATES found by a full-transcript scan and by live detection. Keep each candidate unless the transcript clearly contradicts it, '
            'you may reword for clarity and merge true duplicates, and you should add any further items you find. '
            'If presentation slides are provided, use them as context and cite them as "Slide N" in evidence where relevant. '
            'MIND MAP: "mind_map" describes the meeting. root = meeting title. branches = 3 to 8 main themes. Each theme has children: key points, decisions, '
            'action items (with owner), risks and open questions, as short labels of at most 12 words. Use 2 levels only when the meeting has real structure; '
            'if it is simple, themes can simply have no children. Only use what the transcript supports. '
            'Return ONLY valid JSON with this exact shape: '
            '{"executive_summary":"","meeting_overview":"","key_discussion_points":[""],"decisions":[{"decision":"","evidence":"","timestamp":""}],'
            '"action_items":[{"owner":"","task":"","deadline":"","status":"Assigned","evidence":"","timestamp":""}],'
            '"open_questions":[{"question":"","raised_by":"","timestamp":"","status":"Unanswered","evidence":""}],'
            '"risks":[{"risk":"","evidence":""}],"topics":[{"name":"","summary":"","evidence":""}],'
            '"mind_map":{"root":"","branches":[{"label":"","children":[{"label":"","children":[]}]}]},'
            '"important_moments":[{"timestamp":"","description":"","evidence":""}],"changes_from_previous_meeting":[""],"participants":[""]}. '
            'Keep it concise but comprehensive.')
    user=(f'Meeting: {m["title"]}\nTranscript:\n{transcript}\n\n'
          + (f'PRESENTATION SLIDES:\n{slide_text}\n\n' if slide_text else '')
          + f'{cand_block}\n\nPrevious project context:\n{previous_context(mid)}')
    rep=groq_json(system,user)
    if not isinstance(rep,dict): rep={}

    # Safety net: every candidate survives, even if the model left it out.
    rep['decisions']=merge_items('decision',[x for x in (clean_decision(i) for i in (rep.get('decisions') or [])) if x],cand_dec)
    rep['action_items']=merge_items('task',[x for x in (clean_action(i) for i in (rep.get('action_items') or [])) if x],cand_act)
    rep['open_questions']=merge_items('question',[x for x in (clean_question(i) for i in (rep.get('open_questions') or [])) if x],cand_q)
    print(f'[SessionMind] report {mid}: decisions={len(rep["decisions"])} action_items={len(rep["action_items"])} open_questions={len(rep["open_questions"])}')

    rep.update({'meeting_id':mid,'title':m['title'],'project':p['name'],'started_at':m['started_at'],'ended_at':m['ended_at'],'screenshot_count':len(s),'slides':slides})
    sources=[f'Live capture: {len(t)} transcript lines']
    for mt in mats: sources.append(f"{'Recording' if mt['kind']=='recording' else 'Slides'}: {mt['filename']} ({mt['detail']})")
    if slides and not any(mt['kind']=='slides' for mt in mats): sources.append(f'Slides: {len(slides)} slides')
    rep['sources']=sources

    tree=build_mind_map(rep)
    html_path=DATA/f'report_{mid}.html'; docx_path=DATA/f'report_{mid}.docx'
    render_html(rep,rows,shots,html_path,tree); render_docx(rep,rows,shots,docx_path,BASE,tree)
    c=db(); c.execute('UPDATE meetings SET report=? WHERE id=?',(json.dumps(rep,ensure_ascii=False),mid)); c.execute('INSERT INTO reports(meeting_id,html_path,docx_path,created_at) VALUES(?,?,?,?) ON CONFLICT(meeting_id) DO UPDATE SET html_path=excluded.html_path,docx_path=excluded.docx_path,created_at=excluded.created_at',(mid,str(html_path.relative_to(BASE)),str(docx_path.relative_to(BASE)),now())); c.commit(); c.close()
    return {**rep,'mind_map_tree':tree,'mind_view':mind_view(tree),'files':{'html':f'/api/reports/{mid}/html','docx':f'/api/reports/{mid}/docx'}}

@app.get('/api/reports/{mid}/docx')
def docx_report(mid:int):
    p=DATA/f'report_{mid}.docx'
    if not p.exists():raise HTTPException(404,'Report not found')
    return FileResponse(p,media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',filename=f'sessionmind_{mid}.docx')
@app.get('/api/reports/{mid}/html')
def report_html(mid:int):
    p=DATA/f'report_{mid}.html'
    if not p.exists():raise HTTPException(404,'Report not found')
    return FileResponse(p,media_type='text/html')
@app.get('/api/search')
def search(q:str,project_id:int|None=None):
    c=db(); like=f'%{q}%'; sql='SELECT t.*,m.title FROM transcript t JOIN meetings m ON m.id=t.meeting_id WHERE (t.text LIKE ? OR t.speaker LIKE ?)'; args=[like,like]
    if project_id:sql+=' AND m.project_id=?';args.append(project_id)
    rows=c.execute(sql+' ORDER BY t.timestamp DESC LIMIT 100',args).fetchall();c.close();return [dict(x) for x in rows]
